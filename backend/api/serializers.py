from django.contrib.auth.models import User
from rest_framework import serializers
from .models import Note, Record, DocumentTemplate
from docx import Document

class UserSerializer(serializers.ModelSerializer):
    class Meta:
        model = User
        fields = ["id", "username", "password", "email"]
        extra_kwargs = {"password": {"write_only": True}}

    def create(self, validated_data):
        user = User.objects.create_user(**validated_data)
        return user
    
class NoteSerializer(serializers.ModelSerializer):
    class Meta:
        model = Note
        fields = ["id", "title", "content", "created_at", "author"]
        extra_kwargs = {"author": {"read_only": True}}
    
class RecordSerializer(serializers.ModelSerializer):
    pdf_file_aws = serializers.FileField(required=False, allow_null=True)
    tracking_mail_receipt_aws = serializers.FileField(required=False, allow_null=True)
    return_receipt_file_aws = serializers.FileField(required=False, allow_null=True)

    class Meta:
        model = Record
        fields = [
            'id', 'title', 'tracking_number', 'return_receipt', 'record_type', 'tracking_type', 'status', 'company_name', 'ceo',
            'cfo', 'mailing_address', 'city', 'state', 'zip', 'email', 'phone_number', 'created_by', 'created_at', 'updated_at', 'pdf_file_aws', 'tracking_mail_receipt_aws', 'return_receipt_file_aws', 'hash', 'uuid', 'access_logs'
        ]
        extra_kwargs = {
            'tracking_number': {'required': True},
            'return_receipt': {'required': False},
            'record_type': {'required': False},
            'tracking_type': {'required': False},
            'status': {'required': False},
            'company_name': {'required': False},
            'ceo': {'required': False},
            'cfo': {'required': False},
            'mailing_address': {'required': False},
            'city': {'required': False},
            'state': {'required': False},
            'zip': {'required': False},
            'email': {'required': False},
            'phone_number': {'required': False},
            'pdf_file_aws': {'required': False},
            'tracking_mail_receipt_aws': {'required': False},
            'return_receipt_file_aws': {'required': False},
            'created_by': {'required': False},
            'created_at': {'required': False},
            'updated_at': {'required': False},
            'hash': {'required': False},
            'access_logs': {'required': False},
        }

class DocumentTemplateSerializer(serializers.ModelSerializer):
    class Meta:
        model = DocumentTemplate
        fields = ["id", "name", "description", "file", "placeholders"]
        extra_kwargs = {"file": {"write_only": True}}

    def get_placeholders(self, obj):
        # Logic to extract placeholders from the document
        doc = Document(obj.file.path)
        placeholders = set()
        for paragraph in doc.paragraphs:
            for run in paragraph.runs:
                if '{' in run.text and '}' in run.text:
                    placeholders.add(run.text.strip('{}'))
        return list(placeholders)

class GenerateDocumentsSerializer(serializers.Serializer):
    templates = serializers.ListField(child=serializers.IntegerField())
    placeholders = serializers.DictField(child=serializers.CharField())

